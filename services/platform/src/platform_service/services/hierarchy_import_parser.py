"""Parse hierarchy CSV/XLSX uploads into a desired org-tree graph."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from mc_contracts.errors import ErrorCode
from mc_foundation.problem import AppError
from openpyxl import load_workbook

from platform_service.db.models.hierarchy_user import (
    ROLE_AREA_MANAGER,
    ROLE_PO,
    ROLE_SHASTIYA_KORMI,
)

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 20_000

# Canonical column keys used after header alias resolution.
COL_DIVISION = "division"
COL_DISTRICT = "district"
COL_UPAZILA = "upazila"
COL_AM_NAME = "am_name"
COL_AM_ID = "am_id"
COL_PO_NAME = "po_name"
COL_PO_ID = "po_id"
COL_SK_NAME = "sk_name"
COL_SK_ID = "sk_id"

REQUIRED_COLUMNS = (
    COL_DIVISION,
    COL_DISTRICT,
    COL_UPAZILA,
    COL_AM_NAME,
    COL_AM_ID,
    COL_PO_NAME,
    COL_PO_ID,
    COL_SK_NAME,
    COL_SK_ID,
)

_HEADER_ALIASES: dict[str, str] = {
    "division": COL_DIVISION,
    "district": COL_DISTRICT,
    "upazila": COL_UPAZILA,
    "am name": COL_AM_NAME,
    "am_name": COL_AM_NAME,
    "am mhealth account": COL_AM_ID,
    "am mhealth account id": COL_AM_ID,
    "am_id": COL_AM_ID,
    "am id": COL_AM_ID,
    "rural po name": COL_PO_NAME,
    "po name": COL_PO_NAME,
    "po_name": COL_PO_NAME,
    "po user_id": COL_PO_ID,
    "po user id": COL_PO_ID,
    "po_user_id": COL_PO_ID,
    "po id": COL_PO_ID,
    "po_id": COL_PO_ID,
    "sk name": COL_SK_NAME,
    "sk_name": COL_SK_NAME,
    "sk user_id": COL_SK_ID,
    "sk user id": COL_SK_ID,
    "sk_user_id": COL_SK_ID,
    "sk id": COL_SK_ID,
    "sk_id": COL_SK_ID,
}


@dataclass(frozen=True, slots=True)
class DesiredGeoPath:
    """Division → district → upazila name path (display names, not normalized)."""

    division: str
    district: str
    upazila: str

    @property
    def division_key(self) -> str:
        return _norm_key(self.division)

    @property
    def district_key(self) -> tuple[str, str]:
        return (self.division_key, _norm_key(self.district))

    @property
    def upazila_key(self) -> tuple[str, str, str]:
        return (*self.district_key, _norm_key(self.upazila))


@dataclass
class DesiredUser:
    """One hierarchy user aggregated across all matching spreadsheet rows."""

    user_id: int
    name: str
    role: str
    parent_id: int | None
    division: str
    district: str
    # Normalized (casefolded) upazila names under this user's district.
    upazila_keys: set[str] = field(default_factory=set)


@dataclass
class DesiredHierarchy:
    """Fully validated desired tenant hierarchy from an import file."""

    geo_paths: dict[tuple[str, str, str], DesiredGeoPath] = field(default_factory=dict)
    users: dict[int, DesiredUser] = field(default_factory=dict)

    @property
    def division_names(self) -> dict[str, str]:
        """Normalized division key → display name."""
        out: dict[str, str] = {}
        for path in self.geo_paths.values():
            out.setdefault(path.division_key, path.division)
        return out

    @property
    def district_names(self) -> dict[tuple[str, str], str]:
        out: dict[tuple[str, str], str] = {}
        for path in self.geo_paths.values():
            out.setdefault(path.district_key, path.district)
        return out

    @property
    def upazila_names(self) -> dict[tuple[str, str, str], str]:
        return {key: path.upazila for key, path in self.geo_paths.items()}


def _norm_key(value: str) -> str:
    return value.strip().casefold()


def _normalize_header(raw: str) -> str:
    return " ".join(raw.strip().casefold().split())


def _import_error(detail: str, *, status: int = 400) -> AppError:
    return AppError(
        ErrorCode.HIERARCHY_IMPORT_INVALID.value,
        detail,
        status=status,
    )


def _cell_str(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _parse_positive_int(raw: str, *, field_label: str, row_number: int) -> int:
    text = raw.strip()
    if not text:
        raise _import_error(f"Row {row_number}: {field_label} is required.")
    try:
        # Allow Excel float-looking ids that survived as "422.0"
        value = int(float(text)) if "." in text else int(text)
    except ValueError as exc:
        raise _import_error(f"Row {row_number}: {field_label} must be an integer, got {raw!r}.") from exc
    if value <= 0:
        raise _import_error(f"Row {row_number}: {field_label} must be a positive integer.")
    return value


def _require_nonempty(raw: str, *, field_label: str, row_number: int) -> str:
    text = raw.strip()
    if not text:
        raise _import_error(f"Row {row_number}: {field_label} is required.")
    return text


def _map_headers(raw_headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, header in enumerate(raw_headers):
        if header is None or not str(header).strip():
            continue
        key = _HEADER_ALIASES.get(_normalize_header(str(header)))
        if key is None:
            continue
        if key in mapping:
            raise _import_error(f"Duplicate column for '{key}' in header row.")
        mapping[key] = idx
    missing = [col for col in REQUIRED_COLUMNS if col not in mapping]
    if missing:
        raise _import_error(
            "Missing required columns: " + ", ".join(missing) + ". Expected Division, District, "
            "Upazila, AM name, AM mHealth Account, Rural PO Name, PO User_Id, Sk Name, SK user_id."
        )
    return mapping


def _detect_format(filename: str | None) -> str:
    name = (filename or "").strip().casefold()
    suffix = Path(name).suffix
    if suffix == ".csv":
        return "csv"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    raise _import_error(
        f"Unsupported file type {suffix or '(none)'}; upload a .csv or .xlsx file.",
    )


def _read_capped_bytes(data: bytes) -> bytes:
    if len(data) > MAX_IMPORT_BYTES:
        raise AppError(
            ErrorCode.PAYLOAD_TOO_LARGE.value,
            f"Hierarchy import exceeds {MAX_IMPORT_BYTES} bytes.",
            status=413,
        )
    if not data:
        raise _import_error("Uploaded file is empty.")
    return data


def _rows_from_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    return [[_cell_str(cell) for cell in row] for row in reader]


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    workbook = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise _import_error("Workbook has no active sheet.")
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([_cell_str(cell) for cell in row])
        return rows
    finally:
        workbook.close()


def _merge_user(
    desired: DesiredHierarchy,
    *,
    user_id: int,
    name: str,
    role: str,
    parent_id: int | None,
    division: str,
    district: str,
    upazila_key: str,
    row_number: int,
) -> None:
    existing = desired.users.get(user_id)
    if existing is None:
        desired.users[user_id] = DesiredUser(
            user_id=user_id,
            name=name,
            role=role,
            parent_id=parent_id,
            division=division,
            district=district,
            upazila_keys={upazila_key},
        )
        return

    conflicts: list[str] = []
    if existing.name != name:
        conflicts.append(f"name '{existing.name}' vs '{name}'")
    if existing.role != role:
        conflicts.append(f"role '{existing.role}' vs '{role}'")
    if existing.parent_id != parent_id:
        conflicts.append(f"parent_id {existing.parent_id} vs {parent_id}")
    if _norm_key(existing.division) != _norm_key(division) or _norm_key(existing.district) != _norm_key(
        district
    ):
        conflicts.append(f"district '{existing.division}/{existing.district}' vs '{division}/{district}'")
    if conflicts:
        raise _import_error(
            f"Row {row_number}: user id {user_id} conflicts with earlier rows ({'; '.join(conflicts)})."
        )
    existing.upazila_keys.add(upazila_key)


def build_desired_hierarchy_from_rows(rows: list[list[str]]) -> DesiredHierarchy:
    """Validate spreadsheet rows and return the aggregated desired hierarchy."""
    if not rows:
        raise _import_error("Uploaded file has no rows.")

    header_idx = None
    for idx, row in enumerate(rows):
        if any(cell.strip() for cell in row):
            header_idx = idx
            break
    if header_idx is None:
        raise _import_error("Uploaded file has no header row.")

    header_map = _map_headers(rows[header_idx])
    desired = DesiredHierarchy()
    data_row_count = 0

    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(cell.strip() for cell in row):
            continue
        data_row_count += 1
        if data_row_count > MAX_IMPORT_ROWS:
            raise _import_error(f"Hierarchy import exceeds {MAX_IMPORT_ROWS} data rows.")

        def col(key: str) -> str:
            idx = header_map[key]
            return row[idx] if idx < len(row) else ""

        division = _require_nonempty(col(COL_DIVISION), field_label="Division", row_number=offset)
        district = _require_nonempty(col(COL_DISTRICT), field_label="District", row_number=offset)
        upazila = _require_nonempty(col(COL_UPAZILA), field_label="Upazila", row_number=offset)
        am_name = _require_nonempty(col(COL_AM_NAME), field_label="AM name", row_number=offset)
        am_id = _parse_positive_int(col(COL_AM_ID), field_label="AM mHealth Account", row_number=offset)
        po_name = _require_nonempty(col(COL_PO_NAME), field_label="Rural PO Name", row_number=offset)
        po_id = _parse_positive_int(col(COL_PO_ID), field_label="PO User_Id", row_number=offset)
        sk_name = _require_nonempty(col(COL_SK_NAME), field_label="Sk Name", row_number=offset)
        sk_id = _parse_positive_int(col(COL_SK_ID), field_label="SK user_id", row_number=offset)

        if len({am_id, po_id, sk_id}) < 3:
            raise _import_error(
                f"Row {offset}: AM, PO, and SK user ids must be distinct (got {am_id}, {po_id}, {sk_id})."
            )

        path = DesiredGeoPath(division=division, district=district, upazila=upazila)
        desired.geo_paths.setdefault(path.upazila_key, path)
        upazila_key = path.upazila_key[2]

        _merge_user(
            desired,
            user_id=am_id,
            name=am_name,
            role=ROLE_AREA_MANAGER,
            parent_id=None,
            division=division,
            district=district,
            upazila_key=upazila_key,
            row_number=offset,
        )
        _merge_user(
            desired,
            user_id=po_id,
            name=po_name,
            role=ROLE_PO,
            parent_id=am_id,
            division=division,
            district=district,
            upazila_key=upazila_key,
            row_number=offset,
        )
        _merge_user(
            desired,
            user_id=sk_id,
            name=sk_name,
            role=ROLE_SHASTIYA_KORMI,
            parent_id=po_id,
            division=division,
            district=district,
            upazila_key=upazila_key,
            row_number=offset,
        )

    if data_row_count == 0:
        raise _import_error("Uploaded file has a header but no data rows.")

    # Child upazilas must be ⊆ parent upazilas (same rule as live API).
    for user in desired.users.values():
        if user.parent_id is None:
            continue
        parent = desired.users.get(user.parent_id)
        if parent is None:
            raise _import_error(f"User {user.user_id} references missing parent_id {user.parent_id}.")
        missing = user.upazila_keys - parent.upazila_keys
        if missing:
            raise _import_error(
                f"User {user.user_id} upazilas {sorted(missing)} are not assigned to parent {parent.user_id}."
            )

    return desired


def parse_hierarchy_import_file(*, filename: str | None, data: bytes) -> DesiredHierarchy:
    """Parse capped CSV/XLSX bytes into a validated desired hierarchy."""
    payload = _read_capped_bytes(data)
    fmt = _detect_format(filename)
    rows = _rows_from_csv(payload) if fmt == "csv" else _rows_from_xlsx(payload)
    return build_desired_hierarchy_from_rows(rows)
